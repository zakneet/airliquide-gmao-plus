"""Petite bibliothèque d'assemblage du rapport Word à partir du modèle de l'institut (python-docx)."""
import copy
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PLACEHOLDER = re.compile(r"(\*\*.+?\*\*|`.+?`|\{\{.+?\}\})")


class Rapport:
    kind = "docx"

    def __init__(self, template):
        self.doc = Document(template)
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        self.fig = 0
        self.tab = 0
        self._setup_styles()
        self._setup_page(self.doc.sections[0])
        # Word proposera de mettre à jour les champs (table des matières, listes) à l'ouverture.
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        self.doc.settings.element.append(upd)

    # ---------- mise en page ----------
    def _setup_page(self, s):
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)

    def _setup_styles(self):
        st = self.doc.styles
        n = st["Normal"]
        n.font.name = "Calibri"
        n.font.size = Pt(11)
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.3
        n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size, before, after in (("Heading 1", 18, 0, 14), ("Heading 2", 14, 16, 6), ("Heading 3", 12, 12, 4)):
            s = st[name]
            s.font.name = "Calibri"
            s.font.size = Pt(size)
            s.font.bold = True
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st["Heading 1"].font.color.rgb = RGBColor(0x1F, 0x3A, 0x6E)
        st["Heading 2"].font.color.rgb = RGBColor(0x1F, 0x3A, 0x6E)
        st["Heading 3"].font.color.rgb = RGBColor(0x36, 0x5F, 0x91)
        cap = st["Caption"]
        cap.font.size = Pt(9.5)
        cap.font.italic = False
        cap.font.bold = True
        cap.font.color.rgb = RGBColor(0x1F, 0x3A, 0x6E)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)

    # ---------- texte ----------
    def _runs(self, p, text, size=None, color=None):
        for part in PLACEHOLDER.split(text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 11) - 1)
                r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            elif part.startswith("{{"):
                r = p.add_run("[" + part[2:-2] + "]"); r.bold = True
                hl = OxmlElement("w:highlight"); hl.set(qn("w:val"), "yellow"); r._element.get_or_add_rPr().append(hl)
            else:
                r = p.add_run(part)
            if size and not part.startswith("`"):
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return p

    def p(self, text, style=None, align=None, size=None, space_after=None, keep_next=False):
        para = self.doc.add_paragraph(style=style)
        self._runs(para, text, size=size)
        if align:
            para.alignment = align
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        if keep_next:
            para.paragraph_format.keep_with_next = True
        return para

    def bullets(self, items):
        for t in items:
            para = self.doc.add_paragraph(style="List Bullet")
            para.paragraph_format.space_after = Pt(2)
            self._runs(para, t)

    def numbered(self, items):
        for t in items:
            para = self.doc.add_paragraph(style="List Number")
            para.paragraph_format.space_after = Pt(2)
            self._runs(para, t)

    def h1(self, text, page_break=True):
        para = self.doc.add_paragraph(text, style="Heading 1")
        para.paragraph_format.page_break_before = page_break
        return para

    def h2(self, text):
        return self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text):
        return self.doc.add_paragraph(text, style="Heading 3")

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- champs Word ----------
    def _field(self, para, instr, placeholder=""):
        def fc(t):
            r = para.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._element.append(e)
        fc("begin")
        r = para.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._element.append(it)
        fc("separate")
        para.add_run(placeholder)
        fc("end")

    def toc(self, instr, hint):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self._field(para, instr, hint)

    def caption(self, kind, text, above=False):
        para = self.doc.add_paragraph(style="Caption")
        label = "Figure" if kind == "fig" else "Tableau"
        if kind == "fig":
            self.fig += 1; n = self.fig
        else:
            self.tab += 1; n = self.tab
        para.add_run(f"{label} ")
        self._field(para, f"SEQ {label} \\* ARABIC", str(n))
        para.add_run(f" : {text}")
        if above:
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_after = Pt(4)
            para.paragraph_format.space_before = Pt(8)
        return para

    # ---------- figures et tableaux ----------
    def figure(self, path, caption, width_cm=15.5, max_height_cm=17.5, landscape=False):
        from PIL import Image
        w, h = Image.open(path).size
        width = width_cm
        if width * h / w > max_height_cm:
            width = max_height_cm * w / h
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(str(path), width=Cm(width))
        self.caption("fig", caption)

    def figure_side(self, path, caption, width_cm=6.5):
        """Photo placée à côté du texte en LaTeX ; en Word, simple figure centrée de taille réduite."""
        self.figure(path, caption, width_cm=width_cm, max_height_cm=8)

    def figure_placeholder(self, instruction, caption):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        self._runs(para, "{{" + instruction + "}}")
        self.caption("fig", caption)

    def table(self, header, rows, caption, widths=None, size=9.5):
        self.caption("tab", caption, above=True)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, txt in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            self._runs(c.paragraphs[0], f"**{txt}**", size=size, color=RGBColor(0xFF, 0xFF, 0xFF))
            tcPr = c._element.get_or_add_tcPr()
            sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "1F3A6E")
            tcPr.append(sh)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            for i, txt in enumerate(row):
                cells[i].text = ""
                self._runs(cells[i].paragraphs[0], str(txt), size=size)
                if r_i % 2 == 1:
                    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "EEF3FA")
                    cells[i]._element.get_or_add_tcPr().append(sh)
        for row in t.rows:
            cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().append(cant)
            for c in row.cells:
                for para in c.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    para.paragraph_format.space_after = Pt(2)
                    para.paragraph_format.line_spacing = 1.0
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    # ---------- sections, en-têtes, pieds de page ----------
    def new_section(self, number_format=None, start=None, header_text=None, footer_page=True):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_page(s)
        pg = s._sectPr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType"); s._sectPr.append(pg)
        if number_format:
            pg.set(qn("w:fmt"), number_format)
        if start is not None:
            pg.set(qn("w:start"), str(start))
        s.header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
        hp = s.header.paragraphs[0]
        for r in list(hp.runs):
            r._element.getparent().remove(r._element)
        if header_text:
            r = hp.add_run(header_text); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x59, 0x65, 0x77)
            pPr = hp._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "9AA7BC")):
                b.set(qn(f"w:{k}"), v)
            bd.append(b); pPr.append(bd)
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT  # jc est inséré après pBdr (ordre du schéma respecté)
        fp = s.footer.paragraphs[0]
        for r in list(fp.runs):
            r._element.getparent().remove(r._element)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_page:
            self._field(fp, "PAGE", "1")
        return s

    def first_section_clean(self):
        s = self.doc.sections[0]
        s.header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False

    def save(self, path):
        self.doc.save(str(path))
